import os
import pdfplumber
import docx

folder = r"c:\Users\shrut\Desktop\Loan_verification_system\sampledocs,pdf, and excel"

def inspect_file(filename):
    full_path = os.path.join(folder, filename)
    ext = os.path.splitext(filename)[1].lower()
    print("="*70)
    print(f"FILE: {filename}")
    print("="*70)
    
    if ext == ".pdf":
        with pdfplumber.open(full_path) as p:
            for i, page in enumerate(p.pages[:3]):
                print(f"--- Page {i+1} ---")
                print(page.extract_text()[:1000])
                
    elif ext == ".docx":
        doc = docx.Document(full_path)
        print("Paragraphs:")
        for p in doc.paragraphs[:10]:
            if p.text.strip():
                print("  ", p.text.strip())
        print("Tables:")
        for t in doc.tables[:3]:
            for r in t.rows[:6]:
                row_txt = [c.text.strip().replace('\n', ' ') for c in r.cells if c.text.strip()]
                if row_txt:
                    print("   | " + " | ".join(row_txt[:6]))
                    
    elif ext == ".doc":
        with open(full_path, "rb") as f:
            raw = f.read()
        import re
        strings = re.findall(rb'[\x20-\x7E]{4,}', raw)
        clean_text = " ".join([s.decode('ascii', errors='ignore') for s in strings[:60]])
        print("Text extract:", clean_text[:1200])

for f in sorted(os.listdir(folder)):
    inspect_file(f)
